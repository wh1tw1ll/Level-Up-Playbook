#!/usr/bin/env python3
"""Draft redlined version of DOVA McCarthy CM Precon Agreement."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)
style.paragraph_format.space_after = Pt(4)
style.paragraph_format.line_spacing = 1.15

def add_para(text, bold=False, color=None, alignment=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if bold: run.bold = True
    if color: run.font.color.rgb = color
    if alignment: p.alignment = alignment
    return p

def add_change(change_type, text):
    p = doc.add_paragraph()
    run_label = p.add_run(f'[{change_type}] ')
    run_label.bold = True
    if change_type == 'ADDED':
        run_label.font.color.rgb = RGBColor(0, 128, 0)
    elif change_type == 'MODIFIED':
        run_label.font.color.rgb = RGBColor(0, 0, 180)
    elif change_type == 'DELETED':
        run_label.font.color.rgb = RGBColor(180, 0, 0)
    run_text = p.add_run(text)
    run_text.font.size = Pt(9)
    run_text.font.color.rgb = RGBColor(80, 80, 80)
    return p

def add_heading(text, level=1):
    return doc.add_heading(text, level=level)

# ══════════════════════════════════════════════
add_para('DOVA Rancho Cordova Arena', bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
add_para('Interim Schematic Design Preconstruction Agreement\nand Critical CM Commercial Terms', bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)

add_heading('Parties & Project')
add_para('Owner: KozPure Development, LLC')
add_para('Construction Manager: McCarthy Building Companies, Inc.')
add_para('Project: DOVA Rancho Cordova Arena - Rancho Cordova, California')
add_para('Date: June 26, 2026')
add_para('Consultant: Level Up Project Development, LLC')

# ── 1 ──
add_heading('1. Intent and Current Path')
add_para('The Owner is proceeding with an Owner-direct design team structure and intends to release Perkins & Will to continue schematic design on or around July 6, 2026.')

add_para('McCarthy is being engaged for a limited interim preconstruction role during the schematic design phase to support cost and schedule validation, procurement planning, permitting/AHJ coordination, constructability review, value engineering, and evaluation of the appropriate future construction delivery path.')

add_para('This Agreement does not award McCarthy any construction contract, CMAR agreement, GMP, early work package, or exclusive right to future work. The Owner retains full flexibility to continue with McCarthy, modify McCarthy\'s role, solicit other construction managers, competitively bid the work, procure scopes directly, pursue another delivery strategy, or terminate this Agreement at any time.')

# ── 2 ──
add_heading('2. Term, Fee, and Termination')
add_para('The interim preconstruction term shall begin upon written authorization (the "Authorization Date") and continue through completion of the schematic design phase as defined in the attached Deliverables Schedule (Exhibit A), unless earlier terminated by the Owner.')
add_change('ADDED', 'Defines Authorization Date as term start. References Exhibit A (Deliverables Schedule) for measurable completion rather than open-ended "phase completion."')

add_para('The schematic design phase is currently anticipated to take approximately five (5) to eight (8) weeks from the Authorization Date. This duration is an estimate only and is not guaranteed by the Owner.')

add_para('McCarthy shall provide the interim preconstruction services for the schematic design phase for a lump sum fee of Seventy-Five Thousand Dollars ($75,000).')

add_change('ADDED', 'The $75K fee is supported by a Staffing Budget (Exhibit B) showing roles, hours, and blended rates. This substantiates the fee upfront and prevents claims of understaffing or unpaid overtime.')

add_para('The lump sum fee is intended to cover all of McCarthy\'s preconstruction services through completion of the schematic design phase. No additional fee shall be owed for the schematic design phase unless approved in advance in writing by the Owner.')

add_para('Any proposed change to McCarthy\'s scope or fee shall be documented in a written change order signed by both Parties. McCarthy shall provide a price for requested additional scope within five (5) business days of the Owner\'s written request. The Owner shall have no obligation to authorize any change.')

add_change('ADDED', 'Change order process: Owner direction > McCarthy prices within 5 business days > Owner decides.')

add_para('The Owner may terminate this Agreement at any time for convenience upon written notice. If terminated before SD phase completion, McCarthy shall be paid only the prorated portion of the $75,000 lump sum fee based on the number of calendar days of authorized services actually performed divided by fifty-six (56) calendar days, less any prior payments. No termination fee, lost profit, demobilization, acceleration, future fee, or other compensation shall be owed.')

add_para('Reimbursable expenses require prior written Owner approval. Reimbursable expenses are limited to: (a) travel outside the Sacramento metropolitan area, (b) printing/reproduction exceeding normal practice, and (c) other expenses specifically approved in writing. No markups shall be applied to reimbursable expenses.')
add_change('MODIFIED', 'Reimburables enumerated with categories. Markups prohibited.')

add_para('No construction costs, trade commitments, early procurement, subcontractor awards, general conditions, insurance charges, construction requirements, or other construction-phase costs are authorized unless separately approved in writing by the Owner.')

# ── 3 ──
add_heading('3. Interim Preconstruction Scope and Deliverables')
add_para('During the schematic design phase, McCarthy shall provide normal and customary CM preconstruction services required to advance schematic design and evaluate a potential fast-track construction start, including the services and deliverables listed in Exhibit A. At minimum, McCarthy\'s services shall include:')

items = [
    'Cost estimating and budget validation at the 30%, 60%, and 100% SD milestones',
    'Design-phase cost feedback and trade-off analysis',
    'Facilitated value engineering workshops (minimum two)',
    'Constructability review',
    'Schedule and phasing input with critical path analysis',
    'Early release package planning (structural steel, earthwork, concrete, MEP long-lead)',
    'Long-lead procurement planning with lead-time recommendations',
    'AHJ and permitting support',
    'Site logistics and enabling work planning',
    'Procurement and bid package strategy',
    'Trade market input and outreach (minimum one market conditions report)',
    'Risk identification and mitigation (maintain risk register)',
    'General conditions planning',
    'Weekly progress reporting to Owner and Level Up'
]
for item in items:
    add_para(f'  * {item}')

add_change('ADDED', 'Scope items keyed to SD milestones (30/60/100%). Specific deliverables: VE workshops (2 minimum), risk register, market conditions report.')

add_para('McCarthy\'s work product shall be clear, written, open-book, and sufficient for the Owner to make informed decisions regarding cost, schedule, procurement, delivery structure, and future CM selection.')

add_para('All work product, including estimates, schedules, reports, data, analyses, and other deliverables prepared by McCarthy under this Agreement, shall be the exclusive property of the Owner. The Owner may share such work product with any party, including other construction managers, without restriction or need for further consent.')
add_change('ADDED', 'Work product ownership: all deliverables belong to Owner, can be shared with future CMs.')

add_para('Confidentiality. Neither Party shall disclose confidential information of the other Party (including budget data, fee structures, design intent, business terms, and entitlement strategy) to third parties without prior written consent, except as required by law or to the extent the information is or becomes publicly available through no fault of the receiving Party. This obligation shall survive termination.')
add_change('ADDED', 'Mutual confidentiality clause. Critical because McCarthy sees Owner budget, fees, and project strategy.')

# ── 4 ──
add_heading('4. Potential Future Construction Delivery Options')
add_para('The Owner is evaluating the following potential future construction-phase approaches. The final structure will be determined later by the Owner based on design maturity, budget alignment, market conditions, procurement strategy, funding, schedule, risk allocation, and overall project needs.')

add_para('')
add_para('Option 1: CMAR with GMP - 2.75% CM Fee')
add_para('If the Owner elects to proceed with McCarthy under a CMAR with GMP structure, the applicable CM fee shall be two and three-quarters percent (2.75%) of the GMP amount, excluding Owner-direct scopes, Owner-purchased items, and contingency. The stated CM fee is McCarthy\'s sole compensation for overhead, profit, supervision, management, coordination, and administrative costs. No separate home office overhead, project management markup, profit-on-profit, supervision fee, or similar charge shall be applied beyond the stated fee.')

add_change('MODIFIED', 'Clarified 2.75% excludes Owner-direct scopes and contingency. Fee is SOLE profit compensation - no separate O&P, no stacked markups.')

add_para('')
add_para('Option 2: CMAR Without GMP / CM Agent - 1.50% CM Fee')
add_para('If the Owner elects to proceed with McCarthy under a CMAR without GMP or CM Agent structure, the applicable CM fee shall be one and one-half percent (1.50%), with the same exclusion and compensation principles as Option 1.')

add_para('The Owner may select either option, modify either option, or pursue another delivery path. This Agreement does not obligate the Owner to proceed with McCarthy under either structure.')

# ── 5 ──
add_heading('5. Critical CM Commercial Terms')
add_para('As a condition of this interim engagement and McCarthy\'s continued consideration for a future construction role, McCarthy shall acknowledge and support the following commercial principles:')

principles = [
    ('Open Book', 'All costs shall be fully open book, including fee calculations, staffing, general conditions, construction requirements, insurance, bonds, contingencies, subcontractor pricing, self-perform pricing, allowances, and reimbursable costs. The Owner shall have audit rights.'),
    ('Competitive Buyout', 'The work shall be competitively bought out unless the Owner approves another approach in writing. No subcontractor, supplier, LOI, procurement, or trade commitment may be made without prior written Owner approval.'),
    ('General Conditions', 'Shall be separately identified, detailed by position and duration, open book, and approved by the Owner. The Owner will not accept inflated staffing, duplicative charges, hidden overhead, or costs that should be covered by the CM fee.'),
    ('Construction Requirements', 'Shall be separately identified, open book, and approved by the Owner. Where appropriate, shall be competitively bid or market-priced.'),
    ('No Duplicative Profit Centers', 'The stated CM fee is McCarthy\'s sole compensation. McCarthy shall not recover duplicative compensation through overlapping fee, GCs, CRs, insurance markups, self-perform markups, administrative charges, home office charges, subcontractor markups, or other profit centers.'),
    ('Self-Perform Work', 'Must be identified in advance, competitively validated unless waived, and approved by the Owner. McCarthy shall disclose all cost assumptions including labor, equipment, material, supervision, fee, markup, and contingency.'),
    ('Owner-Direct Scopes', 'The Owner may procure any scope directly. McCarthy shall not receive the full CM fee on Owner-direct scopes unless approved by Owner. If McCarthy provides meaningful coordination support, Owner may consider a coordination fee of up to 0.50% on a scope-by-scope basis.'),
    ('Insurance', 'Structure under review. McCarthy shall clearly identify insurance assumptions, costs, exclusions, deductibles, and markups. During the interim precon term, McCarthy shall maintain: CGL ($2M/$5M), WC statutory, and Professional Liability ($1M) covering preconstruction services. Certificates shall be provided prior to commencing services.'),
    ('Contingency and Savings', 'Shall be separately identified, tracked, and Owner-governed. Use of CM contingency requires Owner approval. Unused contingency, buyout savings, and avoided costs accrue to Owner unless otherwise agreed.'),
    ('Owner Approval Rights', 'Owner retains approval rights over delivery structure, procurement strategy, bid packages, subcontractor awards, self-perform work, Owner-direct scopes, GCs, CRs, insurance, contingency, GMP structure, and all changes.')
]

for title, desc in principles:
    add_para(f'{title}: {desc}', bold=False)
    add_para('')

add_change('ADDED', 'Minimum insurance requirements for precon period: CGL $2M/$5M, PL $1M. Certificates required before starting.')

# ── 6 ──
add_heading('6. General Provisions')

add_para('6.1 Reservation of Rights. Same as original - Owner reserves all rights.')

add_para('6.2 Indemnification. Each Party shall indemnify, defend, and hold harmless the other Party from and against claims, damages, losses, and expenses arising from the indemnifying Party\'s negligent acts or omissions in performing its obligations under this Agreement, to the extent caused by such acts or omissions.')
add_change('ADDED', 'Mutual indemnification for negligence. Important because McCarthy is giving design-phase advice that affects cost/schedule decisions.')

add_para('6.3 Non-Solicitation. During the term and for six (6) months thereafter, McCarthy shall not solicit, recruit, or hire any design consultant, engineer, or other professional retained by the Owner for the Project without the Owner\'s written consent.')
add_change('ADDED', 'Protects P&W, Buehler, Wood Rodgers from being recruited by McCarthy during/after precon.')

add_para('6.4 Governing Law. California. Venue: Sacramento County.')
add_change('ADDED', 'California law, Sacramento venue.')

add_para('6.5 Entire Agreement. This Agreement with exhibits constitutes the entire agreement. No amendment except in writing signed by both Parties.')

add_para('6.6 Notices. In writing, delivered personally, by email with confirmation, or by mail.')

add_para('6.7 No Third-Party Beneficiaries.')

add_para('6.8 Counterparts. May be executed in counterparts and by electronic signature.')

add_change('ADDED', 'Boilerplate sections 6.5-6.8: Entire agreement, notices, no third-party beneficiaries, counterparts.')

# ── Signature ──
add_heading('Accepted and Agreed')
add_para('KozPure Development, LLC', bold=True)
add_para('By: ______________________________')
add_para('Name: ____________________________')
add_para('Title: _____________________________')
add_para('Date: _____________________________')
add_para('')
add_para('McCarthy Building Companies, Inc.', bold=True)
add_para('By: ______________________________')
add_para('Name: ____________________________')
add_para('Title: _____________________________')
add_para('Date: _____________________________')
add_para('')
add_para('Reviewed by:')
add_para('Level Up Project Development, LLC', bold=True)
add_para('By: ______________________________')
add_para('Name: ____________________________')
add_para('Title: _____________________________')
add_para('Date: _____________________________')

# ── Exhibits ──
add_heading('Exhibits')
add_para('Exhibit A - Deliverables Schedule (SD milestones with specific outputs by week)')
add_para('Exhibit B - Staffing Budget / Level of Effort (roles, hours, blended rates)')

# ── Save ──
out_dir = r'C:\Users\HermesAdmin\OneDrive - levelup-pd.com\Documents\00 - LUNA Output\02 - DOVA Arena'
os.makedirs(out_dir, exist_ok=True)
output_path = os.path.join(out_dir, 'DOVA_McCarthy_CM_Precon_Agreement_REDLINED.docx')
doc.save(output_path)
print(f'DONE: {output_path}')
